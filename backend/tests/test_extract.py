import io
import json

import pytest
from docx import Document

from app.ingestion_service.steps.step1_extract import extract


def test_text_and_markdown():
    assert extract("notes.md", b"# Title\nBody")[0]["text"] == "# Title\nBody"


def test_html_drops_scripts_and_styles():
    html = b"<html><style>p{}</style><script>alert(1)</script><p>Hello</p><div>World</div></html>"
    text = extract("page.html", html)[0]["text"]
    assert "Hello" in text and "World" in text
    assert "alert" not in text and "p{}" not in text


def test_json_is_pretty_printed():
    text = extract("data.json", json.dumps({"name": "Asha"}).encode())[0]["text"]
    assert '"name": "Asha"' in text


def test_docx_includes_tables():
    doc = Document()
    doc.add_paragraph("Iron tablets daily.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Dose", "60 mg"
    buffer = io.BytesIO()
    doc.save(buffer)
    text = extract("guide.docx", buffer.getvalue())[0]["text"]
    assert "Iron tablets daily." in text and "Dose | 60 mg" in text


def test_images_return_no_text():
    assert extract("photo.PNG", b"\x89PNG") == []


def test_unsupported_type_raises():
    with pytest.raises(ValueError, match="Unsupported"):
        extract("virus.exe", b"x")
