"""
STEP 1 — EXTRACT: turn a file into plain text.

Output is a list of "sections": {"text": "...", "page": 3}
(page is None for formats without pages). Keeping the page lets answers cite "page 3".

Supported: .pdf .docx .txt .md .csv .json .html .htm
Images (.png .jpg .jpeg .webp) and videos (.mp4 .mov .webm .mpeg) return no text here —
step 2 (caption) turns them into text with the model.

Extend: add OCR for scanned PDFs, audio files, a URL fetcher, etc.
"""
import io
import json
from html.parser import HTMLParser
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

TEXT_TYPES = {".txt", ".md", ".csv"}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
VIDEO_TYPES = {".mp4": "video/mp4", ".mov": "video/mov", ".webm": "video/webm", ".mpeg": "video/mpeg"}
SUPPORTED = {".pdf", ".docx", ".json", ".html", ".htm", *TEXT_TYPES, *IMAGE_TYPES, *VIDEO_TYPES}


def extract(filename: str, data: bytes) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED))}")

    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext in (".html", ".htm"):
        return [{"text": _html_to_text(_decode(data)), "page": None}]
    if ext == ".json":
        return [{"text": json.dumps(json.loads(_decode(data)), indent=2, ensure_ascii=False), "page": None}]
    if ext in TEXT_TYPES:
        return [{"text": _decode(data), "page": None}]
    return []  # images and videos: nothing to extract as text yet, step 2 describes them


def _decode(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> list[dict]:
    reader = PdfReader(io.BytesIO(data))
    return [
        {"text": text, "page": number}
        for number, page in enumerate(reader.pages, start=1)
        if (text := (page.extract_text() or "").strip())
    ]


def _docx(data: bytes) -> list[dict]:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # tables are not part of doc.paragraphs
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [{"text": "\n\n".join(parts), "page": None}]


class _TextOnly(HTMLParser):
    SKIP = {"script", "style", "noscript"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skipping += 1
        elif tag in ("p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skipping:
            self._skipping -= 1

    def handle_data(self, data):
        if not self._skipping:
            self.parts.append(data)


def _html_to_text(html: str) -> str:
    parser = _TextOnly()
    parser.feed(html)
    return "".join(parser.parts)
